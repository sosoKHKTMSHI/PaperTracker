from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Article


HEADERS = [
    "ID", "Autores principais", "Nome curto", "Ano", "Título", "Periódico", "DOI", "PMID", "PMCID",
    "Abstract disponível", "Acesso aberto", "Status do PDF", "Mensagem da verificação", "PDF baixado",
    "Texto estruturado", "Fonte", "Status", "Link principal",
]


def row(article: Article) -> list[str]:
    return [
        article.id,
        article.authors_main,
        article.short_name,
        article.year,
        article.title,
        article.journal,
        article.doi,
        article.pmid,
        article.pmcid,
        article.abstract_status,
        article.oa_status,
        article.pdf_display,
        article.pdf_check_message,
        "Sim" if article.pdf_downloaded else "Não",
        article.structured_display,
        article.source,
        article.status,
        article.main_link,
    ]


def write_xlsx(path: str | Path, articles: list[Article], include_abstracts: bool = False) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Status dos artigos"
    headers = HEADERS + (["Abstract"] if include_abstracts else [])
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    fills = {
        "available": PatternFill("solid", fgColor="C6EFCE"),
        "partial": PatternFill("solid", fgColor="E2F0D9"),
        "missing": PatternFill("solid", fgColor="FFF2CC"),
        "error": PatternFill("solid", fgColor="F4CCCC"),
        "working": PatternFill("solid", fgColor="D9EAF7"),
        "conflict": PatternFill("solid", fgColor="E4DFEC"),
        "pending": PatternFill("solid", fgColor="E7E6E6"),
    }
    for article in articles:
        values = row(article) + ([article.abstract] if include_abstracts else [])
        ws.append(values)
        fill = fills.get(article.color_state, fills["pending"])
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        link_cell = ws.cell(ws.max_row, HEADERS.index("Link principal") + 1)
        if article.main_link:
            link_cell.hyperlink = article.main_link
            link_cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12, 25, 18, 8, 50, 25, 29, 13, 15, 24, 18, 24, 45, 14, 19, 18, 34, 55]
    if include_abstracts:
        widths.append(90)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _summary_counts(articles: list[Article]) -> dict[str, int]:
    return {
        "total": len(articles),
        "confirmed": sum(a.pdf_check_status == "Download confirmado" for a in articles),
        "unverified": sum(bool(a.pdf_url) and a.pdf_check_status == "Não verificado" for a in articles),
        "blocked": sum(a.pdf_check_status == "Acesso bloqueado" for a in articles),
        "invalid": sum(a.pdf_check_status == "Link inválido" for a in articles),
        "temporary": sum(a.pdf_check_status == "Erro temporário" for a in articles),
        "missing": sum(not a.pdf_url for a in articles),
        "structured": sum(bool(a.json_url or a.xml_url) for a in articles),
        "abstracts": sum(a.abstract_status != "Não" for a in articles),
        "downloaded": sum(bool(a.pdf_downloaded) for a in articles),
    }


def write_txt(path: str | Path, articles: list[Article], include_abstracts: bool = False) -> Path:
    path = Path(path)
    counts = _summary_counts(articles)
    lines = [
        "RELATÓRIO DE RASTREAMENTO DE ARTIGOS",
        f"Gerado em: {datetime.now().astimezone().strftime('%d/%m/%Y %H:%M:%S %z')}",
        "",
        "RESUMO",
        f"Artigos cadastrados: {counts['total']}",
        f"PDFs com download confirmado: {counts['confirmed']}",
        f"Links de PDF ainda não verificados: {counts['unverified']}",
        f"PDFs com acesso bloqueado: {counts['blocked']}",
        f"Links de PDF inválidos: {counts['invalid']}",
        f"Erros temporários de verificação: {counts['temporary']}",
        f"PDFs não localizados: {counts['missing']}",
        f"PDFs baixados nesta sessão/projeto: {counts['downloaded']}",
        f"Textos estruturados disponíveis: {counts['structured']}",
        f"Abstracts disponíveis: {counts['abstracts']}",
        "",
    ]
    for article in articles:
        lines.extend([
            "=" * 78,
            f"{article.id} — {article.authors_main or 'Autores não informados'} — {article.short_name or 'Sem nome curto'}",
            "=" * 78,
            "",
            f"Título: {article.title or 'Não informado'}",
            f"Periódico: {article.journal or 'Não informado'}",
            f"Ano: {article.year or 'Não informado'}",
            f"DOI: {article.doi or 'Não localizado'}",
            f"PMID: {article.pmid or 'Não localizado'}",
            f"PMCID: {article.pmcid or 'Não localizado'}",
            f"Abstract disponível: {article.abstract_status}",
            f"Acesso aberto: {article.oa_status}",
            f"PDF: {article.pdf_display}",
            f"Mensagem da verificação: {article.pdf_check_message or 'Não informado'}",
            f"PDF baixado: {'Sim' if article.pdf_downloaded else 'Não'}",
            f"Erro no download: {article.pdf_download_error or 'Não'}",
            f"JSON BioC: {'Localizado' if article.json_url else 'Não localizado'}",
            f"XML BioC: {'Localizado' if article.xml_url else 'Não localizado'}",
            f"Fonte principal: {article.source or 'Não informada'}",
            f"Status: {article.status}",
            f"Link principal: {article.main_link or 'Não localizado'}",
            f"Link PDF: {article.pdf_url or 'Não localizado'}",
            "",
        ])
        if include_abstracts:
            lines.extend(["ABSTRACT", article.abstract or "Não disponível", ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_html(path: str | Path, articles: list[Article]) -> Path:
    path = Path(path)
    counts = _summary_counts(articles)
    rows = []
    colors = {
        "available": "#c6efce", "partial": "#e2f0d9", "missing": "#fff2cc", "error": "#f4cccc",
        "working": "#d9eaf7", "conflict": "#e4dfec", "pending": "#e7e6e6",
    }
    for article in articles:
        values = row(article)
        cells = []
        for index, value in enumerate(values):
            if HEADERS[index] == "Link principal" and value:
                cells.append(f'<td><a href="{escape(value, quote=True)}">{escape(value)}</a></td>')
            else:
                cells.append(f"<td>{escape(str(value))}</td>")
        rows.append(f'<tr style="background:{colors.get(article.color_state, colors["pending"])}">' + "".join(cells) + "</tr>")
    summary = (
        f"Artigos: {counts['total']} | PDFs confirmados: {counts['confirmed']} | PDFs não localizados: {counts['missing']} | "
        f"Textos estruturados: {counts['structured']} | PDFs baixados: {counts['downloaded']}"
    )
    html = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Status dos artigos</title>
<style>
body{{font-family:Arial,sans-serif;margin:20px;color:#222}}button{{padding:8px 14px;margin-bottom:12px}}
table{{border-collapse:collapse;width:100%;font-size:11px}}th,td{{border:1px solid #aaa;padding:5px;vertical-align:top}}
th{{background:#1f4e78;color:white;position:sticky;top:0}}a{{color:#0645ad;word-break:break-all}}
@media print{{button{{display:none}}th{{position:static}}body{{margin:0}}}}
</style></head><body><button onclick="window.print()">Imprimir</button>
<h1>Status dos artigos</h1><p>Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}</p><p>{escape(summary)}</p>
<table><thead><tr>{''.join(f'<th>{escape(h)}</th>' for h in HEADERS)}</tr></thead><tbody>{''.join(rows)}</tbody></table>
</body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding="utf-8")
    return path
